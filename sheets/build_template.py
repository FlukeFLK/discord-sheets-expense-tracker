#!/usr/bin/env python3
"""
สร้างไฟล์ Excel/Google Sheets "บัญชีรายรับ รายจ่าย เงินออม" (4 ชีต) ด้วย openpyxl

    pip install -r requirements.txt
    python build_template.py income-expense-tracker.xlsx

ตั้ง PREVIEW=1 เพื่อใช้ฟอนต์ไทย (Loma) สำหรับเรนเดอร์ภาพตัวอย่างด้วย LibreOffice
ข้อมูลตัวอย่างในหน้า Transaction สุ่มจาก seed คงที่ (ไม่ใช่ข้อมูลจริง)
"""
import random
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, PieChart, Reference
from openpyxl.chart.data_source import NumFmt
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.drawing.text import (CharacterProperties, Font as DFont, Paragraph,
                                   ParagraphProperties, RegularTextRun)
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else "tracker.xlsx"

# ---------------------------------------------------------------- palette
INK = "1E293B"      # primary text
INK2 = "475569"     # secondary text
MUTED = "94A3B8"
LINE = "E2E8F0"
HEAD = "334155"     # neutral table header
INPUT = "FEF9C3"    # pale yellow = cell you fill / choose
SYS = "F1F5F9"      # system / calculated area

TYPES = [
    dict(key="inc", name="รายรับ", icon="💰", setup_col="B", n=15,
         color="1BAF7A", dark="0E7A55", pale="E8F7F1"),
    dict(key="exp", name="รายจ่าย", icon="💸", setup_col="D", n=25,
         color="EB6834", dark="B4461C", pale="FDEEE7"),
    dict(key="sav", name="เงินออม", icon="🏦", setup_col="F", n=15,
         color="2A78D6", dark="1C5CAB", pale="EAF2FC"),
]
T = {t["key"]: t for t in TYPES}
CASH = dict(color="4A3AA7", dark="3B2E8A", pale="EEECF8")
RATE = dict(color="EDA100", dark="8A5D00", pale="FDF5E1")
CNT = dict(color="64748B", dark="475569", pale="F1F5F9")
# categorical slots for pie slices (validated order) + neutral "other"
CAT = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "4A3AA7"]
OTHER = "B8B6AE"

import os
PREVIEW = os.environ.get("PREVIEW") == "1"
FONT = "Loma" if PREVIEW else "Arial"
NUM = '#,##0;[Red]-#,##0;"-"'
NUM2 = '#,##0.00;[Red]-#,##0.00;"-"'
BAHT = '฿#,##0;[Red]-฿#,##0;฿0'
PCT = '0.0%;[Red]-0.0%;"-"'
DELTA_PCT = '▲ 0.0%;▼ 0.0%;0.0%'
DELTA_NUM = '+#,##0;[Red]-#,##0;"-"'

MONTHS_FULL = ["มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
               "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"]
MONTHS_SHORT = ["ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.",
                "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]

CATEGORIES = {
    "inc": ["เงินเดือน", "โบนัส", "ค่าล่วงเวลา (OT)", "ฟรีแลนซ์/งานเสริม",
            "ขายของออนไลน์", "ดอกเบี้ย/ปันผล", "เงินคืนภาษี", "รายรับอื่นๆ"],
    "exp": ["ค่าอาหาร", "ค่าเดินทาง", "ค่าที่พัก/ค่าเช่า", "ค่าน้ำ-ค่าไฟ",
            "ค่าโทรศัพท์/เน็ต", "ของใช้ในบ้าน", "ช้อปปิ้ง", "สุขภาพ/ยา",
            "ประกันภัย", "สังสรรค์/บันเทิง", "การศึกษา", "ผ่อนชำระ/หนี้",
            "ให้ครอบครัว", "ท่องเที่ยว", "Subscription", "ทำบุญ/บริจาค",
            "รายจ่ายอื่นๆ"],
    "sav": ["เงินฝากออมทรัพย์", "เงินสำรองฉุกเฉิน", "กองทุน SSF/RMF", "หุ้น/ETF",
            "ออมทอง", "ออมเพื่อเกษียณ", "ออมอื่นๆ"],
}

# Transaction table geometry
TX_FIRST, TX_LAST = 6, 5005
# whole-column references: summaries read every row, no upper limit
R_DATE = "Transaction!$B:$B"
R_TYPE = "Transaction!$C:$C"
R_CAT = "Transaction!$D:$D"
R_AMT = "Transaction!$E:$E"
SETUP_SHORT = "Setup!$J$6:$J$17"
SETUP_FULL = "Setup!$I$6:$I$17"
SETUP_YEARS = "Setup!$H$6:$H$17"
STACK_CAT = "Setup!$L$6:$L$60"
STACK_TYPE = "Setup!$M$6:$M$60"

# ---------------------------------------------------------------- style helpers
thin = Side(style="thin", color=LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
BOTTOM = Border(bottom=thin)
WHITE_SIDE = Side(style="medium", color="FFFFFF")
CARD_BORDER = Border(left=WHITE_SIDE, right=WHITE_SIDE, top=WHITE_SIDE, bottom=WHITE_SIDE)


def font(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def al(h="left", v="center", wrap=False, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)


def cells(ws, rng):
    if ":" not in rng:
        yield ws[rng]
        return
    for row in ws[rng]:
        for c in row:
            yield c


def style(ws, rng, f=None, fl=None, b=None, a=None, nf=None):
    for c in cells(ws, rng):
        if f is not None:
            c.font = f
        if fl is not None:
            c.fill = fl
        if b is not None:
            c.border = b
        if a is not None:
            c.alignment = a
        if nf is not None:
            c.number_format = nf


def put(ws, ref, value, f=None, fl=None, b=None, a=None, nf=None, merge=None):
    """Write value to ref (anchor); optionally merge ref:merge and style the whole block."""
    rng = f"{ref}:{merge}" if merge else ref
    if merge:
        ws.merge_cells(rng)
    ws[ref] = value
    style(ws, rng, f, fl, b, a, nf)
    return ws[ref]


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def heights(ws, spec):
    for r, h in spec.items():
        ws.row_dimensions[r].height = h


def title_bar(ws, last_col, text, sub_formula, sub_last_col):
    put(ws, "B2", text, font(18, True, "FFFFFF"), fill(INK), a=al("left", indent=1),
        merge=f"{last_col}2")
    put(ws, "B3", sub_formula, font(10, False, INK2), a=al("left", indent=1),
        merge=f"{sub_last_col}3")
    heights(ws, {1: 8, 2: 38, 3: 26, 4: 8})


def section(ws, ref, merge, text, color=INK):
    put(ws, ref, text, font(12, True, color), b=Border(bottom=Side(style="medium", color=color)),
        a=al("left"), merge=merge)


def card(ws, c1, c2, label, value, value_fmt, sub, sub_fmt, theme, row=5):
    put(ws, f"{c1}{row}", label, font(10, True, "FFFFFF"), fill(theme["dark"]), CARD_BORDER,
        al("left", indent=1), merge=f"{c2}{row}")
    put(ws, f"{c1}{row+1}", value, font(18, True, INK), fill(theme["pale"]), CARD_BORDER,
        al("left", indent=1), value_fmt, merge=f"{c2}{row+1}")
    put(ws, f"{c1}{row+2}", sub, font(9, False, INK2), fill(theme["pale"]), CARD_BORDER,
        al("left", indent=1), sub_fmt, merge=f"{c2}{row+2}")


def header_row(ws, row, labels, color=HEAD, fcolor="FFFFFF", h="center"):
    """labels: list of (col_or_range, text)."""
    for col, text in labels:
        if ":" in col:
            a, b = col.split(":")
            put(ws, f"{a}{row}", text, font(10, True, fcolor), fill(color), BORDER, al(h, wrap=True),
                merge=f"{b}{row}")
        else:
            put(ws, f"{col}{row}", text, font(10, True, fcolor), fill(color), BORDER, al(h, wrap=True))


# ---------------------------------------------------------------- chart helpers
def cp(sz=900, b=False, color=INK2):
    return CharacterProperties(sz=sz, b=b, solidFill=color,
                               latin=DFont(typeface=FONT), cs=DFont(typeface="Tahoma"))


def txpr(sz=900, b=False, color=INK2):
    p = Paragraph(pPr=ParagraphProperties(defRPr=cp(sz, b, color)), endParaRPr=cp(sz, b, color))
    return RichText(p=[p])


def chart_title(text, sz=1100):
    c = cp(sz, True, INK)
    p = Paragraph(pPr=ParagraphProperties(defRPr=c), r=[RegularTextRun(rPr=c, t=text)])
    return Title(tx=Text(rich=RichText(p=[p])), overlay=False)


def frame(chart):
    chart.graphical_properties = GraphicalProperties(
        solidFill="FFFFFF", ln=LineProperties(solidFill=LINE, w=9525))
    chart.roundedCorners = False


def place(ws, chart, col1, row1, col2, row2, colw):
    """Anchor chart from top-left of (col1,row1) to bottom-right of (col2,row2), 1-based cells."""
    pad = 5 * 9525
    w_last = int(colw[gl(col2)] * 7 + 5)
    a = TwoCellAnchor()
    a._from = AnchorMarker(col=col1 - 1, colOff=pad, row=row1 - 1, rowOff=pad)
    a.to = AnchorMarker(col=col2 - 1, colOff=max(0, w_last * 9525 - pad), row=row2, rowOff=0)
    chart.anchor = a
    ws.add_chart(chart)


def bar_chart(title, legend_pos="b"):
    ch = BarChart()
    ch.type = "col"
    ch.grouping = "clustered"
    ch.gapWidth = 70
    ch.overlap = -8
    ch.title = chart_title(title)
    ch.style = 2
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.x_axis.txPr = txpr(900)
    ch.y_axis.txPr = txpr(800, color=MUTED)
    ch.y_axis.numFmt = NumFmt(formatCode="#,##0", sourceLinked=False)
    from openpyxl.chart.axis import ChartLines
    ch.y_axis.majorGridlines = ChartLines()
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="EEF2F6", w=6350))
    ch.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill="CBD5E1", w=9525))
    ch.x_axis.majorTickMark = "none"
    ch.y_axis.majorTickMark = "none"
    ch.legend.position = legend_pos
    ch.legend.txPr = txpr(900)
    frame(ch)
    return ch


def color_series(s, color):
    s.graphicalProperties = GraphicalProperties(solidFill=color, ln=LineProperties(solidFill=color))


def slice_colors(series, colors):
    for i, c in enumerate(colors):
        pt = DataPoint(idx=i, spPr=GraphicalProperties(
            solidFill=c, ln=LineProperties(solidFill="FFFFFF", w=19050)))
        series.dPt.append(pt)


def doughnut(ws, title, cats_ref, vals_ref, colors, legend=True):
    ch = DoughnutChart(holeSize=58)
    ch.title = chart_title(title)
    ch.add_data(vals_ref, titles_from_data=False)
    ch.set_categories(cats_ref)
    slice_colors(ch.series[0], colors)
    dl = DataLabelList()
    dl.showPercent = True
    dl.showVal = dl.showCatName = dl.showSerName = dl.showLegendKey = False
    dl.numFmt = '[>=0.04]0%;""'
    dl.txPr = txpr(800, True, INK)
    ch.dataLabels = dl
    if legend:
        ch.legend.position = "r"
        ch.legend.txPr = txpr(850)
    else:
        ch.legend = None
    frame(ch)
    return ch


def pie_labeled(title, cats_ref, vals_ref, colors):
    ch = PieChart()
    ch.title = chart_title(title)
    ch.add_data(vals_ref, titles_from_data=False)
    ch.set_categories(cats_ref)
    slice_colors(ch.series[0], colors)
    dl = DataLabelList()
    dl.showCatName = True
    dl.showPercent = True
    dl.showVal = dl.showSerName = dl.showLegendKey = False
    dl.showLeaderLines = True
    dl.separator = "\n"
    dl.position = "outEnd"
    dl.numFmt = "0%"
    dl.txPr = txpr(800, False, INK)
    ch.dataLabels = dl
    ch.legend = None
    frame(ch)
    return ch


# ================================================================ workbook
wb = Workbook()
wb.calculation = CalcProperties(fullCalcOnLoad=True)
ws_dash = wb.active
ws_dash.title = "Dashboard Overview"
ws_mon = wb.create_sheet("Monthly Overview")
ws_tx = wb.create_sheet("Transaction")
ws_set = wb.create_sheet("Setup")

for ws, tab in [(ws_dash, "2A78D6"), (ws_mon, "1BAF7A"), (ws_tx, "EB6834"), (ws_set, "64748B")]:
    ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    ws.sheet_format.defaultRowHeight = 18
    ws.sheet_format.customHeight = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

# ================================================================ SETUP
s = ws_set
widths(s, {"A": 2, "B": 30, "C": 3, "D": 30, "E": 3, "F": 30, "G": 5,
           "H": 10, "I": 14, "J": 9, "K": 3, "L": 24, "M": 11})
put(s, "B2", "🧾 Setup — ตั้งค่าหมวดหมู่รายรับ รายจ่าย เงินออม", font(18, True, "FFFFFF"),
    fill(INK), a=al("left", indent=1), merge="F2")
put(s, "B3", "พิมพ์ชื่อหมวดหมู่ในช่องสีเหลือง เพิ่มหมวดใหม่ต่อท้ายได้เลย (ห้ามเว้นแถวว่างคั่นกลาง) "
    "· ระบบจะนำไปทำ Drop-down ในหน้า Transaction และสรุปใน Dashboard อัตโนมัติ",
    font(10, False, INK2), a=al("left", wrap=True, indent=1), merge="F3")
heights(s, {1: 8, 2: 38, 3: 34, 4: 20, 5: 30})

for t in TYPES:
    col = t["setup_col"]
    first, last = 6, 5 + t["n"]
    put(s, f"{col}5",
        f'="{t["icon"]} หมวดหมู่{t["name"]}  ("&COUNTA({col}{first}:{col}{last})&"/{t["n"]})"',
        font(11, True, "FFFFFF"), fill(t["dark"]), BORDER, al("left", indent=1))
    for i in range(t["n"]):
        r = first + i
        c = s[f"{col}{r}"]
        c.value = CATEGORIES[t["key"]][i] if i < len(CATEGORIES[t["key"]]) else None
        c.font = font(10, False, INK)
        c.fill = fill(INPUT)
        c.border = BORDER
        c.alignment = al("left", indent=1)

# system lists
put(s, "H4", "⚙️ ข้อมูลระบบ (ไม่ต้องแก้ไข)", font(9, True, INK2), merge="M4")
header_row(s, 5, [("H", "ปี (ค.ศ.)"), ("I", "เดือน"), ("J", "ย่อ"),
                  ("L", "รายการหมวดหมู่รวม\n(ใช้ทำ Drop-down)"), ("M", "ประเภท")],
           color="64748B")
for i in range(12):
    r = 6 + i
    for col, v in (("H", 2024 + i), ("I", MONTHS_FULL[i]), ("J", MONTHS_SHORT[i])):
        put(s, f"{col}{r}", v, font(10, False, INK2), fill(SYS), BORDER, al("center"))

nI, nE, nS = "COUNTA($B$6:$B$20)", "COUNTA($D$6:$D$30)", "COUNTA($F$6:$F$20)"
for k in range(1, 56):
    r = 5 + k
    s[f"L{r}"] = (f"=IF({k}<={nI},INDEX($B$6:$B$20,{k}),"
                  f"IF({k}<={nI}+{nE},INDEX($D$6:$D$30,{k}-{nI}),"
                  f"IF({k}<={nI}+{nE}+{nS},INDEX($F$6:$F$20,{k}-{nI}-{nE}),\"\")))")
    s[f"M{r}"] = (f'=IF({k}<={nI},"รายรับ",IF({k}<={nI}+{nE},"รายจ่าย",'
                  f'IF({k}<={nI}+{nE}+{nS},"เงินออม","")))')
    style(s, f"L{r}:M{r}", font(9, False, INK2), fill(SYS), BORDER, al("left", indent=1))

# how-to box
guide = [
    ("📘 วิธีใช้งาน", True),
    ("1) Setup — แก้/เพิ่มหมวดหมู่ในช่องสีเหลือง (สูงสุด: รายรับ 15 · รายจ่าย 25 · ออม 15 หมวด)", False),
    ("2) Transaction — บันทึกทุกรายการ: วันที่ → ประเภท → หมวดหมู่ → จำนวนเงิน → รายละเอียด", False),
    ("3) Dashboard Overview — เลือกปีที่มุมขวาบน ดูภาพรวมทั้งปี ค่าเฉลี่ยรายเดือน Top 10 และ Heat Map", False),
    ("4) Monthly Overview — เลือกปีและเดือน ดูสรุปเดือนนั้นพร้อมเทียบกับเดือนก่อน", False),
    ("5) ข้อมูลในหน้า Transaction เป็นตัวอย่าง ลบได้ทันที (ลบเฉพาะคอลัมน์ B–F ห้ามลบคอลัมน์ 'ตรวจสอบ')", False),
    ("6) ถ้าเปลี่ยนชื่อหมวดหมู่ รายการเก่าที่ใช้ชื่อเดิมจะขึ้น ⚠ ในหน้า Transaction ให้แก้ชื่อให้ตรงกัน", False),
    ("7) แถวไม่พอ: เพิ่มแถวท้ายชีต แล้วคัดลอกแถวสุดท้าย (B–G) วางลงแถวใหม่ — ระบบนับทุกแถวไม่จำกัด", False),
    ("💡 คงเหลือ (Cash Flow) = รายรับ − รายจ่าย − เงินออม   ·   อัตราการออม = เงินออม ÷ รายรับ", False),
    ("🟨 ช่องสีเหลือง = ช่องที่กรอกหรือเลือกได้   ·   ช่องสีเทา = ระบบคำนวณให้ ไม่ต้องแก้ไข", False),
]
g0 = 33
for i, (text, bold) in enumerate(guide):
    r = g0 + i
    put(s, f"B{r}", text, font(11 if bold else 10, bold, INK if bold else INK2),
        fill("F8FAFC"), a=al("left", indent=1), merge=f"F{r}")
    s.row_dimensions[r].height = 24 if bold else 20
s.freeze_panes = "A6"

# ================================================================ TRANSACTION
x = ws_tx
widths(x, {"A": 2, "B": 14, "C": 12, "D": 24, "E": 16, "F": 42, "G": 24})
put(x, "B2", "🧾 Transaction — บันทึกรายรับ รายจ่าย เงินออม", font(18, True, "FFFFFF"),
    fill(INK), a=al("left", indent=1), merge="G2")
put(x, "B3", "กรอกทีละแถว: วันที่ → ประเภท (Drop-down) → หมวดหมู่ (Drop-down) → จำนวนเงิน → รายละเอียด "
    "· คอลัมน์ 'ตรวจสอบ' ระบบเช็กให้อัตโนมัติ · ข้อมูลตัวอย่างลบได้ทันที (ลบเฉพาะคอลัมน์ B–F) "
    "· แถวไม่พอ: เพิ่มแถวท้ายชีต แล้วคัดลอกแถวสุดท้าย (B–G) วางลงไป",
    font(10, False, INK2), a=al("left", wrap=True, indent=1), merge="G3")
put(x, "B4", '="📌 บันทึกแล้ว "&TEXT(COUNTIF($C:$C,"รายรับ")+COUNTIF($C:$C,"รายจ่าย")+COUNTIF($C:$C,"เงินออม"),"#,##0")'
    '&" รายการ     ⚠ ต้องแก้ไข "&COUNTIF($G:$G,"⚠*")&" รายการ     (ระบบนับทุกแถว ไม่จำกัดจำนวน)"',
    font(10, True, INK), fill("F8FAFC"), a=al("left", indent=1), merge="G4")
heights(x, {1: 8, 2: 38, 3: 44, 4: 24, 5: 28})
header_row(x, 5, [("B", "วันที่"), ("C", "ประเภท"), ("D", "หมวดหมู่"),
                  ("E", "จำนวนเงิน (บาท)"), ("F", "รายละเอียด"), ("G", "ตรวจสอบ")], color=INK)

f_cell = font(10, False, INK)
for r in range(TX_FIRST, TX_LAST + 1):
    x[f"B{r}"].number_format = "dd/mm/yyyy"
    x[f"B{r}"].alignment = al("center")
    x[f"C{r}"].alignment = al("center")
    x[f"D{r}"].alignment = al("left", indent=1)
    x[f"E{r}"].number_format = NUM2
    x[f"E{r}"].alignment = al("right", indent=1)
    x[f"F{r}"].alignment = al("left", indent=1)
    x[f"G{r}"] = (f'=IF(COUNTA(B{r}:E{r})=0,"",IF(COUNTA(B{r}:E{r})<4,"⚠ กรอกไม่ครบ",'
                  f'IF(NOT(ISNUMBER(B{r})),"⚠ วันที่ไม่ถูกต้อง",'
                  f'IF(COUNTIFS({STACK_CAT},D{r},{STACK_TYPE},C{r})=0,"⚠ หมวดไม่ตรงประเภท","✓"))))')
    x[f"G{r}"].alignment = al("center")
    x[f"G{r}"].fill = fill("F8FAFC")
    for col in "BCDEFG":
        c = x[f"{col}{r}"]
        c.font = f_cell
        c.border = BOTTOM

# conditional formats: type colours
data_c = f"C{TX_FIRST}:C{TX_LAST}"
data_e = f"E{TX_FIRST}:E{TX_LAST}"
for t in TYPES:
    rule_f = f'$C{TX_FIRST}="{t["name"]}"'
    x.conditional_formatting.add(data_c, FormulaRule(formula=[rule_f], font=Font(color=t["dark"], bold=True),
                                                     fill=fill(t["pale"])))
    x.conditional_formatting.add(data_e, FormulaRule(formula=[rule_f], font=Font(color=t["dark"])))
x.conditional_formatting.add(f"G{TX_FIRST}:G{TX_LAST}", FormulaRule(
    formula=[f'LEFT($G{TX_FIRST},1)="⚠"'], font=Font(color="B91C1C", bold=True), fill=fill("FEE2E2")))
x.conditional_formatting.add(f"G{TX_FIRST}:G{TX_LAST}", FormulaRule(
    formula=[f'$G{TX_FIRST}="✓"'], font=Font(color="15803D", bold=True)))

# data validation
dv_date = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                         showErrorMessage=True, errorTitle="วันที่ไม่ถูกต้อง",
                         error="กรุณากรอกวันที่ เช่น 25/09/2026")
dv_type = DataValidation(type="list", formula1='"รายรับ,รายจ่าย,เงินออม"', allow_blank=True,
                         showErrorMessage=True, errorTitle="เลือกประเภท",
                         error="เลือกจาก Drop-down: รายรับ / รายจ่าย / เงินออม")
dv_cat = DataValidation(type="list", formula1=STACK_CAT, allow_blank=True, showErrorMessage=True,
                        errorTitle="ไม่พบหมวดหมู่นี้",
                        error="เลือกจาก Drop-down หรือเพิ่มหมวดหมู่ใหม่ที่หน้า Setup ก่อน")
dv_amt = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        showErrorMessage=True, errorTitle="จำนวนเงินไม่ถูกต้อง",
                        error="กรอกเป็นตัวเลขบวก (ไม่ต้องใส่เครื่องหมายลบ)")
for dv, col in ((dv_date, "B"), (dv_type, "C"), (dv_cat, "D"), (dv_amt, "E")):
    dv.add(f"{col}{TX_FIRST}:{col}{TX_LAST}")
    x.add_data_validation(dv)
x.freeze_panes = "A6"
x.auto_filter.ref = f"B5:G{TX_LAST}"

# sample data -------------------------------------------------------------
rnd = random.Random(2026)


def rr(a, b, step=10):
    return rnd.randrange(a, b + 1, step)


rows = []


def add(m, d, typ, cat, amt, desc):
    rows.append((date(2026, m, d), typ, cat, float(amt), desc))


I, E, S = "รายรับ", "รายจ่าย", "เงินออม"
for m in range(1, 10):
    add(m, 25, I, "เงินเดือน", 42000, "เงินเดือนประจำเดือน")
    if m in (1, 2, 4, 6, 7, 9):
        add(m, 25, I, "ค่าล่วงเวลา (OT)", rr(1500, 4800, 100), "ค่า OT ทำงานวันหยุด")
    if m in (2, 3, 5, 7, 8, 9):
        add(m, rr(8, 20, 1), I, "ฟรีแลนซ์/งานเสริม", rr(3000, 12000, 500),
            rnd.choice(["งานออกแบบโลโก้", "รับเขียนคอนเทนต์", "งานแปลเอกสาร", "ถ่ายรูปสินค้า"]))
    if m in (3, 6, 8, 9):
        add(m, rr(12, 22, 1), I, "ขายของออนไลน์", rr(800, 3500, 50), "ขายของมือสองใน Marketplace")
    if m in (3, 6, 9):
        add(m, 30, I, "ดอกเบี้ย/ปันผล", rr(350, 900, 10), "ดอกเบี้ยบัญชีออมทรัพย์")
    if m == 1:
        add(1, 15, I, "โบนัส", 63000, "โบนัสประจำปี")
    if m == 5:
        add(5, 12, I, "เงินคืนภาษี", 4860, "ภาษีคืนจากสรรพากร")
        add(5, 20, I, "ดอกเบี้ย/ปันผล", 2400, "เงินปันผลหุ้น")

    add(m, 1, E, "ค่าที่พัก/ค่าเช่า", 8500, "ค่าเช่าคอนโด")
    add(m, 5, E, "ค่าน้ำ-ค่าไฟ", rr(1350, 1750) if m in (4, 5) else rr(850, 1250), "ค่าไฟ + ค่าน้ำ")
    add(m, 10, E, "ค่าโทรศัพท์/เน็ต", 599, "ค่าโทรศัพท์รายเดือน")
    add(m, 12, E, "ค่าโทรศัพท์/เน็ต", 590, "อินเทอร์เน็ตบ้าน")
    add(m, 15, E, "ผ่อนชำระ/หนี้", 4200, "ผ่อนโน้ตบุ๊ก 0% (10 งวด)")
    add(m, 26, E, "ให้ครอบครัว", 5000, "ส่งเงินให้พ่อแม่")
    add(m, 8, E, "Subscription", 419, "Netflix")
    add(m, 20, E, "Subscription", 149, "Spotify")
    for wk, d in enumerate((3, 10, 17, 24), 1):
        add(m, d, E, "ค่าอาหาร", rr(1200, 2000), f"ค่าอาหาร-กาแฟ สัปดาห์ที่ {wk}")
    add(m, rr(6, 28, 1), E, "ค่าอาหาร", rr(600, 1600), rnd.choice(["มื้อพิเศษกับครอบครัว", "ปิ้งย่างกับเพื่อน", "ซื้อของเข้าตู้เย็น"]))
    add(m, 2, E, "ค่าเดินทาง", rr(500, 1000, 100), "เติมเงินบัตร BTS/MRT")
    add(m, rr(9, 18, 1), E, "ค่าเดินทาง", rr(150, 450), "Grab / วินมอเตอร์ไซค์")
    add(m, rr(19, 28, 1), E, "ค่าเดินทาง", rr(600, 1200), "เติมน้ำมัน")
    add(m, rr(4, 14, 1), E, "ของใช้ในบ้าน", rr(300, 1100), "ของใช้ในบ้าน/ห้องน้ำ")
    add(m, rr(5, 27, 1), E, "ช้อปปิ้ง", rr(450, 2500), rnd.choice(["เสื้อผ้า", "รองเท้า", "ของแต่งห้อง", "อุปกรณ์ไอที"]))
    if m in (3, 6, 9):
        add(m, 9 if m == 9 else 15, E, "ช้อปปิ้ง", rr(900, 2800), "โปรเซลล์กลางเดือน")
        add(m, 18, E, "ประกันภัย", 3200, "ประกันสุขภาพ (รายไตรมาส)")
    if m in (2, 5, 7, 9):
        add(m, rr(6, 24, 1), E, "สุขภาพ/ยา", rr(250, 1800), rnd.choice(["ค่ายา/ร้านขายยา", "ทำฟัน", "ตรวจสุขภาพ"]))
    add(m, rr(10, 28, 1), E, "สังสรรค์/บันเทิง", rr(400, 1800), rnd.choice(["ดูหนัง", "สังสรรค์กับเพื่อน", "คอนเสิร์ต", "คาราโอเกะ"]))
    if m in (1, 4, 7, 8):
        add(m, rr(5, 28, 1), E, "ทำบุญ/บริจาค", rr(100, 1000, 50), "ทำบุญ/บริจาค")
    if m == 2:
        add(2, 14, E, "การศึกษา", 1990, "คอร์สออนไลน์ Data Analysis")
    if m == 7:
        add(7, 9, E, "การศึกษา", 650, "ซื้อหนังสือ")
    if m == 4:
        add(4, 12, E, "ท่องเที่ยว", 12800, "เที่ยวสงกรานต์ เชียงใหม่")
    if m == 8:
        add(8, 16, E, "ท่องเที่ยว", 6500, "เที่ยวหัวหิน")

    add(m, 26, S, "เงินฝากออมทรัพย์", 3000, "ฝากออมประจำเดือน")
    add(m, 26, S, "เงินสำรองฉุกเฉิน", 2000, "เติมเงินสำรองฉุกเฉิน")
    add(m, 27, S, "กองทุน SSF/RMF", 2500, "DCA กองทุนลดหย่อนภาษี")
    if m in (2, 3, 5, 6, 8):
        add(m, 27, S, "หุ้น/ETF", rr(2000, 5000, 500), "DCA ETF")
    if m in (1, 4, 7, 9):
        add(m, 28, S, "ออมทอง", rr(1000, 2000, 500), "ออมทองรายเดือน")
    if m == 1:
        add(1, 16, S, "เงินสำรองฉุกเฉิน", 20000, "แบ่งจากโบนัส")
        add(1, 16, S, "หุ้น/ETF", 15000, "ลงทุนจากโบนัส")

rows.sort(key=lambda r: (r[0], {"รายรับ": 0, "รายจ่าย": 1, "เงินออม": 2}[r[1]]))
for i, (d, typ, cat, amt, desc) in enumerate(rows):
    r = TX_FIRST + i
    x[f"B{r}"], x[f"C{r}"], x[f"D{r}"], x[f"E{r}"], x[f"F{r}"] = d, typ, cat, amt, desc
print("sample rows:", len(rows))

# ================================================================ shared blocks
DASH_W = {"A": 2, "B": 26, **{gl(c): 10.5 for c in range(3, 18)}, "R": 4}
TOP_BLOCKS = [  # (type key, cat first col, cat last col, amount col, pct col, bar col)
    ("inc", "B", "B", "C", "D", "E"),
    ("exp", "G", "H", "I", "J", "K"),
    ("sav", "M", "N", "O", "P", "Q"),
]


def top10(ws, row0, sources, period_txt):
    """Top-10 tables (title row0, header row0+1, ranks row0+2..+11, others +12, total +13).
    sources[key] = (cat_range, amount_range, rank_range, total_ref)."""
    for key, c1, c2, ca, cpct, cbar in TOP_BLOCKS:
        t = T[key]
        cat_rng, amt_rng, rank_rng, total_ref = sources[key]
        put(ws, f"{c1}{row0}", f'🏆 Top 10 {t["name"]}{period_txt}', font(11, True, "FFFFFF"),
            fill(t["dark"]), BORDER, al("left", indent=1), merge=f"{cbar}{row0}")
        hdr = [(f"{c1}:{c2}" if c1 != c2 else c1, "หมวดหมู่"), (ca, "จำนวนเงิน"),
               (cpct, "%"), (cbar, "สัดส่วน")]
        header_row(ws, row0 + 1, hdr, color=t["pale"], fcolor=t["dark"])
        first = row0 + 2
        last = row0 + 11
        tot = row0 + 13
        for k in range(1, 11):
            r = row0 + 1 + k
            match = f"MATCH({k},{rank_rng},0)"
            put(ws, f"{c1}{r}", f'=IFERROR("{k}. "&INDEX({cat_rng},{match}),"")',
                font(10, k <= 3, INK), b=BORDER, a=al("left", indent=1),
                merge=(f"{c2}{r}" if c1 != c2 else None))
            put(ws, f"{ca}{r}", f"=IFERROR(INDEX({amt_rng},{match}),\"\")",
                font(10, k <= 3, INK), b=BORDER, a=al("right", indent=1), nf=NUM)
            put(ws, f"{cpct}{r}", f'=IF(OR({ca}{r}="",N({ca}${tot})=0),"",{ca}{r}/{ca}${tot})',
                font(9, False, INK2), b=BORDER, a=al("center"), nf='0.0%')
            put(ws, f"{cbar}{r}",
                f'=IF({ca}{r}="","",REPT("█",MAX(1,ROUND({ca}{r}/MAX(${ca}${first}:${ca}${last})*10,0))))',
                Font(name=FONT, size=7, color=t["color"]), b=BORDER, a=al("left", indent=0))
        r = row0 + 12
        put(ws, f"{c1}{r}", "หมวดอื่นๆ ที่เหลือ", font(9, False, INK2, True), b=BORDER,
            a=al("left", indent=1), merge=(f"{c2}{r}" if c1 != c2 else None))
        put(ws, f"{ca}{r}", f"=MAX(0,{total_ref}-SUM({ca}{first}:{ca}{last}))", font(9, False, INK2),
            b=BORDER, a=al("right", indent=1), nf=NUM)
        put(ws, f"{cpct}{r}", f'=IF(N({ca}${tot})=0,"",{ca}{r}/{ca}${tot})', font(9, False, INK2),
            b=BORDER, a=al("center"), nf='0.0%')
        put(ws, f"{cbar}{r}", None, b=BORDER)
        put(ws, f"{c1}{tot}", f"รวม{t['name']}", font(10, True, t["dark"]), fill(t["pale"]), BORDER,
            al("left", indent=1), merge=(f"{c2}{tot}" if c1 != c2 else None))
        put(ws, f"{ca}{tot}", f"={total_ref}", font(10, True, t["dark"]), fill(t["pale"]), BORDER,
            al("right", indent=1), NUM)
        put(ws, f"{cpct}{tot}", f'=IF(N({ca}{tot})=0,"",1)', font(9, True, t["dark"]), fill(t["pale"]),
            BORDER, al("center"), '0%')
        put(ws, f"{cbar}{tot}", None, fl=fill(t["pale"]), b=BORDER)
        for rr_ in range(row0 + 1, tot + 1):
            ws.row_dimensions[rr_].height = 19


def chart_data_block(ws, row0, sources, col_pairs):
    """Top-7 + 'อื่นๆ' per type (8 rows) for pie / doughnut charts, in the system area."""
    out = {}
    for key, (cc, vc) in zip(("inc", "exp", "sav"), col_pairs):
        t = T[key]
        cat_rng, amt_rng, rank_rng, total_ref = sources[key]
        put(ws, f"{cc}{row0}", f"กราฟ: {t['name']}", font(9, True, INK2), fill("E2E8F0"), BORDER,
            al("left", indent=1))
        put(ws, f"{vc}{row0}", "จำนวนเงิน", font(9, True, INK2), fill("E2E8F0"), BORDER, al("center"))
        for k in range(1, 8):
            r = row0 + k
            match = f"MATCH({k},{rank_rng},0)"
            put(ws, f"{cc}{r}", f'=IFERROR(INDEX({cat_rng},{match}),"")', font(9, False, INK2),
                fill(SYS), BORDER, al("left", indent=1))
            put(ws, f"{vc}{r}", f'=IFERROR(INDEX({amt_rng},{match}),"")', font(9, False, INK2),
                fill(SYS), BORDER, al("right", indent=1), NUM)
        r = row0 + 8
        put(ws, f"{cc}{r}", "อื่นๆ", font(9, False, INK2), fill(SYS), BORDER, al("left", indent=1))
        put(ws, f"{vc}{r}", f'=IF(MAX(0,{total_ref}-SUM({vc}{row0+1}:{vc}{row0+7}))=0,"",'
                            f'MAX(0,{total_ref}-SUM({vc}{row0+1}:{vc}{row0+7})))',
            font(9, False, INK2), fill(SYS), BORDER, al("right", indent=1), NUM)
        out[key] = (ci(cc), ci(vc), row0 + 1, row0 + 8)
    return out


# ================================================================ DASHBOARD OVERVIEW
d = ws_dash
widths(d, DASH_W)
widths(d, {"S": 20, "T": 11, "U": 3, "V": 20, "W": 11, "X": 3, "Y": 20, "Z": 11})
YR = "$P$3"
title_bar(d, "Q", "📊 Dashboard Overview — ภาพรวมการเงินรายปี",
          f'="ข้อมูลปี ค.ศ. "&{YR}&" (พ.ศ. "&({YR}+543)&")   ·   ระบบสรุปอัตโนมัติจากหน้า Transaction"', "L")
put(d, "M3", "เลือกปี (ค.ศ.) ▸", font(10, True, INK), a=al("right"), merge="O3")
put(d, "P3", 2026, font(14, True, INK), fill(INPUT), Border(left=Side(style="thin", color="EAB308"),
    right=Side(style="thin", color="EAB308"), top=Side(style="thin", color="EAB308"),
    bottom=Side(style="thin", color="EAB308")), al("center"), "0", merge="Q3")
dv_y = DataValidation(type="list", formula1=SETUP_YEARS, allow_blank=False, showErrorMessage=True,
                      error="เลือกปีจาก Drop-down")
dv_y.add("P3")
d.add_data_validation(dv_y)

# monthly summary table (rows 10-24) --------------------------------------
MS_H, MS_1, MS_12, MS_TOT, MS_AVG = 10, 11, 22, 23, 24
section(d, "B9", "H9", "📅 สรุปรายเดือน")
header_row(d, MS_H, [("B", "เดือน"), ("C", "รายรับ"), ("D", "รายจ่าย"), ("E", "เงินออม"),
                     ("F", "คงเหลือ"), ("G", "อัตราออม"), ("H", "คงเหลือสะสม")])
for col, t in (("C", T["inc"]), ("D", T["exp"]), ("E", T["sav"]), ("F", CASH)):
    d[f"{col}{MS_H}"].fill = fill(t["dark"])
for i in range(12):
    r = MS_1 + i
    mexpr = f"MATCH($B{r},{SETUP_SHORT},0)"
    put(d, f"B{r}", MONTHS_SHORT[i], font(10, True, INK), b=BORDER, a=al("center"))
    for col, t in (("C", "รายรับ"), ("D", "รายจ่าย"), ("E", "เงินออม")):
        put(d, f"{col}{r}", f'=SUMIFS({R_AMT},{R_TYPE},"{t}",{R_DATE},">="&DATE({YR},{mexpr},1),'
                            f'{R_DATE},"<"&DATE({YR},{mexpr}+1,1))',
            font(10), b=BORDER, a=al("right", indent=1), nf=NUM)
    put(d, f"F{r}", f"=C{r}-D{r}-E{r}", font(10, True, CASH["dark"]), b=BORDER, a=al("right", indent=1), nf=NUM)
    put(d, f"G{r}", f'=IF(C{r}=0,"",E{r}/C{r})', font(10, False, INK2), b=BORDER, a=al("center"), nf=PCT)
    put(d, f"H{r}", f'=IF(C{r}+D{r}+E{r}=0,"",SUM($F${MS_1}:F{r}))', font(10, False, INK2), b=BORDER,
        a=al("right", indent=1), nf=NUM)
put(d, f"B{MS_TOT}", "รวมทั้งปี", font(10, True, "FFFFFF"), fill(HEAD), BORDER, al("center"))
put(d, f"B{MS_AVG}", "เฉลี่ย/เดือน", font(10, True, INK), fill("E2E8F0"), BORDER, al("center"))
for col in "CDEF":
    put(d, f"{col}{MS_TOT}", f"=SUM({col}{MS_1}:{col}{MS_12})", font(10, True, "FFFFFF"), fill(HEAD),
        BORDER, al("right", indent=1), NUM)
    put(d, f"{col}{MS_AVG}", f"=IF(N($P$6)=0,0,{col}{MS_TOT}/$P$6)", font(10, True, INK),
        fill("E2E8F0"), BORDER, al("right", indent=1), NUM)
put(d, f"G{MS_TOT}", f'=IF(C{MS_TOT}=0,"",E{MS_TOT}/C{MS_TOT})', font(10, True, "FFFFFF"), fill(HEAD),
    BORDER, al("center"), PCT)
put(d, f"G{MS_AVG}", f'=IF(C{MS_AVG}=0,"",E{MS_AVG}/C{MS_AVG})', font(10, True, INK), fill("E2E8F0"),
    BORDER, al("center"), PCT)
put(d, f"H{MS_TOT}", None, fl=fill(HEAD), b=BORDER)
put(d, f"H{MS_AVG}", None, fl=fill("E2E8F0"), b=BORDER)
put(d, "B25", "* คงเหลือ = รายรับ − รายจ่าย − เงินออม   ·   ค่าเฉลี่ยคิดจากเดือนที่มีการบันทึกข้อมูล",
    font(8, False, MUTED, True), merge="H25")

# KPI cards (rows 5-7) -----------------------------------------------------
card(d, "B", "C", "💰 รายรับรวม", f"=$C${MS_TOT}", BAHT, f"=$C${MS_AVG}", '"เฉลี่ย/เดือน ฿"#,##0', T["inc"])
card(d, "D", "F", "💸 รายจ่ายรวม", f"=$D${MS_TOT}", BAHT, f"=$D${MS_AVG}", '"เฉลี่ย/เดือน ฿"#,##0', T["exp"])
card(d, "G", "I", "🏦 เงินออมรวม", f"=$E${MS_TOT}", BAHT, f"=$E${MS_AVG}", '"เฉลี่ย/เดือน ฿"#,##0', T["sav"])
card(d, "J", "L", "💵 คงเหลือ (Cash Flow)", f"=$F${MS_TOT}", BAHT, f"=$F${MS_AVG}",
     '"เฉลี่ย/เดือน ฿"#,##0;"เฉลี่ย/เดือน -฿"#,##0;"เฉลี่ย/เดือน ฿0"', CASH)
card(d, "M", "O", "📈 อัตราการออม", f"=IF($C${MS_TOT}=0,0,$E${MS_TOT}/$C${MS_TOT})", "0.0%",
     f'=IF($C${MS_TOT}=0,"ยังไม่มีรายรับในปีนี้","ใช้จ่าย "&TEXT($D${MS_TOT}/$C${MS_TOT},"0%")&" ของรายรับ")',
     None, RATE)
card(d, "P", "Q", "🗓️ เดือนที่มีข้อมูล",
     f"=SUMPRODUCT(--(($C${MS_1}:$C${MS_12}+$D${MS_1}:$D${MS_12}+$E${MS_1}:$E${MS_12})>0))",
     '0" เดือน"',
     f'=TEXT(COUNTIFS({R_DATE},">="&DATE({YR},1,1),{R_DATE},"<"&DATE({YR}+1,1,1)),"#,##0")&" รายการ"',
     None, CNT)
heights(d, {5: 24, 6: 38, 7: 22, 8: 14, 9: 26, MS_H: 24, 25: 16, 26: 10})

# heat maps -----------------------------------------------------------------
HM = {}  # key -> (header row, first, last, total row)
row = 62
for key in ("inc", "exp", "sav"):
    t = T[key]
    hr, first = row, row + 1
    last = first + t["n"] - 1
    tot = last + 1
    HM[key] = (hr, first, last, tot)
    hdr = [("B", f'{t["icon"]} {t["name"]}')] + [(gl(3 + i), MONTHS_SHORT[i]) for i in range(12)] + \
          [("O", "รวมทั้งปี"), ("P", "%"), ("Q", "อันดับ")]
    header_row(d, hr, hdr, color=t["dark"])
    d[f"B{hr}"].alignment = al("left", indent=1)
    d.row_dimensions[hr].height = 24
    src_col = t["setup_col"]
    for i in range(t["n"]):
        r = first + i
        srow = 6 + i
        put(d, f"B{r}", f'=IF(Setup!${src_col}${srow}="","",Setup!${src_col}${srow})', font(10, False, INK),
            b=BORDER, a=al("left", indent=1))
        for m in range(12):
            col = gl(3 + m)
            mexpr = f"MATCH({col}${hr},{SETUP_SHORT},0)"
            put(d, f"{col}{r}",
                f'=IF($B{r}="","",SUMIFS({R_AMT},{R_TYPE},"{t["name"]}",{R_CAT},$B{r},'
                f'{R_DATE},">="&DATE({YR},{mexpr},1),{R_DATE},"<"&DATE({YR},{mexpr}+1,1)))',
                font(9, False, INK), b=BORDER, a=al("center"), nf=NUM)
        put(d, f"O{r}", f'=IF($B{r}="","",SUM(C{r}:N{r}))', font(10, True, INK), b=BORDER,
            a=al("right", indent=1), nf=NUM)
        put(d, f"P{r}", f'=IF(OR($B{r}="",N($O${tot})=0),"",O{r}/$O${tot})', font(9, False, INK2), b=BORDER,
            a=al("center"), nf=PCT)
        put(d, f"Q{r}", f'=IF(OR($B{r}="",N(O{r})<=0),"",COUNTIF($O${first}:$O${last},">"&O{r})'
                        f'+COUNTIF($O${first}:O{r},O{r}))', font(9, True, t["dark"]), b=BORDER,
            a=al("center"), nf="0")
    put(d, f"B{tot}", f"รวม{t['name']}", font(10, True, t["dark"]), fill(t["pale"]), BORDER, al("left", indent=1))
    for m in range(13):
        col = gl(3 + m)
        put(d, f"{col}{tot}", f"=SUM({col}{first}:{col}{last})", font(10, True, t["dark"]), fill(t["pale"]),
            BORDER, al("center" if m < 12 else "right", indent=0 if m < 12 else 1), NUM)
    put(d, f"P{tot}", f'=IF(N(O{tot})=0,"",1)', font(9, True, t["dark"]), fill(t["pale"]), BORDER,
        al("center"), "0%")
    put(d, f"Q{tot}", None, fl=fill(t["pale"]), b=BORDER)
    d.conditional_formatting.add(f"C{first}:N{last}", ColorScaleRule(
        start_type="num", start_value=0, start_color="FFFFFF", end_type="max", end_color=t["color"]))
    row = tot + 2
cash_row = row
put(d, f"B{cash_row}", "💵 คงเหลือสุทธิ", font(10, True, "FFFFFF"), fill(CASH["dark"]), BORDER, al("left", indent=1))
for m in range(13):
    col = gl(3 + m)
    put(d, f"{col}{cash_row}",
        f"={col}{HM['inc'][3]}-{col}{HM['exp'][3]}-{col}{HM['sav'][3]}",
        font(10, True, CASH["dark"]), fill(CASH["pale"]), BORDER, al("center" if m < 12 else "right", indent=0 if m < 12 else 1), NUM)
put(d, f"P{cash_row}", None, fl=fill(CASH["pale"]), b=BORDER)
put(d, f"Q{cash_row}", None, fl=fill(CASH["pale"]), b=BORDER)
d.row_dimensions[cash_row].height = 24
section(d, "B60", "Q60", "🔥 Heat Map — รายละเอียดแยกตามหมวดหมู่ × รายเดือน")
put(d, "B61", "สีเข้ม = จำนวนเงินมาก   ·   สีอ่อน = จำนวนเงินน้อย   ·   แสดงข้อมูลของปีที่เลือกด้านบน   ·   "
    "อันดับ = เรียงจากยอดรวมทั้งปีมากไปน้อย", font(9, False, INK2, True), merge="Q61")
heights(d, {59: 10, 60: 26, 61: 18})

# Top 10 (rows 45-58) -------------------------------------------------------
dash_src = {k: (f"$B${HM[k][1]}:$B${HM[k][2]}", f"$O${HM[k][1]}:$O${HM[k][2]}",
                f"$Q${HM[k][1]}:$Q${HM[k][2]}", f"$O${HM[k][3]}") for k in HM}
section(d, "B27", "Q27", "🍩 สัดส่วนตามหมวดหมู่ & Top 10 (ทั้งปี)")
top10(d, 45, dash_src, "")
heights(d, {26: 10, 27: 26, 44: 10, 45: 24, 46: 22})

# chart data (system area) --------------------------------------------------
put(d, "S27", "⚙️ ข้อมูลประกอบกราฟ (ระบบคำนวณ — ไม่ต้องแก้ไข)", font(9, True, INK2), merge="Z27")
dash_cd = chart_data_block(d, 28, dash_src, [("S", "T"), ("V", "W"), ("Y", "Z")])

# charts ---------------------------------------------------------------------
ch = bar_chart("รายรับ · รายจ่าย · เงินออม รายเดือน")
data = Reference(d, min_col=3, max_col=5, min_row=MS_H, max_row=MS_12)
ch.add_data(data, titles_from_data=True)
ch.set_categories(Reference(d, min_col=2, min_row=MS_1, max_row=MS_12))
for s_, t in zip(ch.series, TYPES):
    color_series(s_, t["color"])
place(d, ch, ci("I"), 9, ci("Q"), 25, DASH_W)

pie_cols = CAT + [OTHER]
for (key, c1, *_), (cc, vc, r1, r2), (a, b) in zip(TOP_BLOCKS, [dash_cd[k] for k in ("inc", "exp", "sav")],
                                                  [("B", "E"), ("G", "K"), ("M", "Q")]):
    dn = doughnut(d, f"สัดส่วน{T[key]['name']}ตามหมวดหมู่",
                  Reference(d, min_col=cc, min_row=r1, max_row=r2),
                  Reference(d, min_col=vc, min_row=r1, max_row=r2), pie_cols)
    place(d, dn, ci(a), 28, ci(b), 43, {**DASH_W})
d.freeze_panes = "A4"
d.print_area = f"A1:Q{cash_row}"

# ================================================================ MONTHLY OVERVIEW
mo = ws_mon
widths(mo, DASH_W)
widths(mo, {"S": 20, "T": 12, "U": 8, "V": 3, "W": 20, "X": 12, "Y": 8, "Z": 3, "AA": 20, "AB": 12, "AC": 8})
MY, MM = "$L$3", "$O$3"
title_bar(mo, "Q", "📈 Monthly Overview — สรุปรายเดือน",
          f'="ข้อมูลเดือน "&{MM}&" "&{MY}&" (พ.ศ. "&({MY}+543)&")   ·   เปรียบเทียบกับเดือนก่อนอัตโนมัติ"', "J")
put(mo, "K3", "ปี ค.ศ. ▸", font(10, True, INK), a=al("right"))
yb = Side(style="thin", color="EAB308")
put(mo, "L3", 2026, font(13, True, INK), fill(INPUT), Border(left=yb, right=yb, top=yb, bottom=yb), al("center"), "0")
put(mo, "M3", "เดือน ▸", font(10, True, INK), a=al("right"), merge="N3")
put(mo, "O3", "กันยายน", font(13, True, INK), fill(INPUT), Border(left=yb, right=yb, top=yb, bottom=yb),
    al("center"), merge="Q3")
dv_my = DataValidation(type="list", formula1=SETUP_YEARS, allow_blank=False, showErrorMessage=True,
                       error="เลือกปีจาก Drop-down")
dv_mm = DataValidation(type="list", formula1=SETUP_FULL, allow_blank=False, showErrorMessage=True,
                       error="เลือกเดือนจาก Drop-down")
dv_my.add("L3")
dv_mm.add("O3")
mo.add_data_validation(dv_my)
mo.add_data_validation(dv_mm)

# system parameters
put(mo, "S2", "⚙️ ส่วนคำนวณของระบบ (ไม่ต้องแก้ไข)", font(9, True, INK2), merge="AC2")
params = [("เดือนที่ (ตัวเลข)", f"=MATCH({MM},{SETUP_FULL},0)", "0"),
          ("วันแรกของเดือน", f"=DATE({MY},$T$3,1)", "dd/mm/yyyy"),
          ("วันแรกเดือนถัดไป", f"=DATE({MY},$T$3+1,1)", "dd/mm/yyyy"),
          ("วันแรกเดือนก่อน", f"=DATE({MY},$T$3-1,1)", "dd/mm/yyyy"),
          ("จำนวนวันในเดือน", "=DAY($T$5-1)", "0")]
for i, (lab, f_, nf_) in enumerate(params):
    r = 3 + i
    put(mo, f"S{r}", lab, font(9, False, INK2), fill(SYS), BORDER, al("left", indent=1))
    put(mo, f"T{r}", f_, font(9, False, INK2), fill(SYS), BORDER, al("center"), nf_)
CUR = f'{R_DATE},">="&$T$4,{R_DATE},"<"&$T$5'
PRV = f'{R_DATE},">="&$T$6,{R_DATE},"<"&$T$4'

# sorting helper: category, amount this month, rank
MH = {}
for key, (cc, vc, rc) in zip(("inc", "exp", "sav"), [("S", "T", "U"), ("W", "X", "Y"), ("AA", "AB", "AC")]):
    t = T[key]
    header_row(mo, 10, [(cc, f"หมวด{t['name']}"), (vc, "ยอดเดือนนี้"), (rc, "อันดับ")], color="64748B")
    first, last = 11, 10 + t["n"]
    for i in range(t["n"]):
        r = first + i
        srow = 6 + i
        sc = t["setup_col"]
        put(mo, f"{cc}{r}", f'=IF(Setup!${sc}${srow}="","",Setup!${sc}${srow})', font(9, False, INK2),
            fill(SYS), BORDER, al("left", indent=1))
        put(mo, f"{vc}{r}", f'=IF({cc}{r}="","",SUMIFS({R_AMT},{R_TYPE},"{t["name"]}",{R_CAT},{cc}{r},{CUR}))',
            font(9, False, INK2), fill(SYS), BORDER, al("right", indent=1), NUM)
        put(mo, f"{rc}{r}", f'=IF(OR({cc}{r}="",N({vc}{r})<=0),"",COUNTIF(${vc}${first}:${vc}${last},">"&{vc}{r})'
                           f'+COUNTIF(${vc}${first}:{vc}{r},{vc}{r}))', font(9, False, INK2), fill(SYS), BORDER,
            al("center"), "0")
    tot = last + 1
    put(mo, f"{cc}{tot}", "รวม", font(9, True, INK2), fill("E2E8F0"), BORDER, al("left", indent=1))
    put(mo, f"{vc}{tot}", f"=SUM({vc}{first}:{vc}{last})", font(9, True, INK2), fill("E2E8F0"), BORDER,
        al("right", indent=1), NUM)
    put(mo, f"{rc}{tot}", None, fl=fill("E2E8F0"), b=BORDER)
    MH[key] = (f"${cc}${first}:${cc}${last}", f"${vc}${first}:${vc}${last}",
               f"${rc}${first}:${rc}${last}", f"${vc}${tot}")

# comparison table (rows 10-14)
section(mo, "B9", "Q9", "📊 ภาพรวมเดือนนี้ เทียบกับเดือนก่อน")
header_row(mo, 10, [("B", "รายการ"), ("C", "เดือนนี้"), ("D", "เดือนก่อน"), ("E", "เปลี่ยนแปลง"), ("F", "%")])
comp = [("inc", "รายรับ"), ("exp", "รายจ่าย"), ("sav", "เงินออม")]
for i, (key, lab) in enumerate(comp):
    r = 11 + i
    t = T[key]
    put(mo, f"B{r}", lab, font(10, True, t["dark"]), fill(t["pale"]), BORDER, al("left", indent=1))
    put(mo, f"C{r}", f'=SUMIFS({R_AMT},{R_TYPE},"{t["name"]}",{CUR})', font(10, True), b=BORDER,
        a=al("right", indent=1), nf=NUM)
    put(mo, f"D{r}", f'=SUMIFS({R_AMT},{R_TYPE},"{t["name"]}",{PRV})', font(10, False, INK2), b=BORDER,
        a=al("right", indent=1), nf=NUM)
r = 14
put(mo, f"B{r}", "คงเหลือ", font(10, True, CASH["dark"]), fill(CASH["pale"]), BORDER, al("left", indent=1))
put(mo, f"C{r}", "=C11-C12-C13", font(10, True), b=BORDER, a=al("right", indent=1), nf=NUM)
put(mo, f"D{r}", "=D11-D12-D13", font(10, False, INK2), b=BORDER, a=al("right", indent=1), nf=NUM)
for r in range(11, 15):
    put(mo, f"E{r}", f"=C{r}-D{r}", font(10, False, INK2), b=BORDER, a=al("right", indent=1), nf=DELTA_NUM)
    put(mo, f"F{r}", f'=IF(D{r}=0,"",(C{r}-D{r})/ABS(D{r}))', font(9, False, INK2), b=BORDER,
        a=al("center"), nf=DELTA_PCT)

# allocation table (rows 16-20)
header_row(mo, 16, [("B", "การใช้รายรับเดือนนี้"), ("C", "จำนวนเงิน"), ("D", "สัดส่วน")])
alloc = [("รายจ่าย", "=C12", T["exp"]), ("เงินออม", "=C13", T["sav"]),
         ("คงเหลือ", "=MAX(0,C14)", CASH)]
for i, (lab, f_, th) in enumerate(alloc):
    r = 17 + i
    put(mo, f"B{r}", lab, font(10, True, th["dark"]), fill(th["pale"]), BORDER, al("left", indent=1))
    put(mo, f"C{r}", f_, font(10), b=BORDER, a=al("right", indent=1), nf=NUM)
    put(mo, f"D{r}", f'=IF($C$11=0,"",C{r}/$C$11)', font(10, False, INK2), b=BORDER, a=al("center"), nf=PCT)
put(mo, "B20", '=IF(C11+C12+C13=0,"ยังไม่มีข้อมูลในเดือนนี้",IF(C14<0,"⚠ เดือนนี้ใช้เงินเกินรายรับ ฿"&TEXT(-C14,"#,##0"),'
    '"✅ เดือนนี้มีเงินเหลือ ฿"&TEXT(C14,"#,##0")))', font(10, True, INK), fill("F8FAFC"), BORDER,
    al("left", indent=1), merge="F20")
mo.conditional_formatting.add("B20", FormulaRule(formula=['$C$14<0'], font=Font(color="B91C1C", bold=True),
                                                 fill=fill("FEE2E2")))
mo.conditional_formatting.add("B20", FormulaRule(formula=['$C$14>0'], font=Font(color="15803D", bold=True),
                                                 fill=fill("DCFCE7")))

# insights (rows 22-25)
put(mo, "B22", "📌 ข้อมูลน่ารู้", font(10, True, "FFFFFF"), fill(HEAD), BORDER, al("left", indent=1), merge="F22")
ins = [("รายจ่ายเฉลี่ยต่อวัน", "=IF(N($T$7)=0,0,$C$12/$T$7)", BAHT),
       ("รายจ่ายก้อนใหญ่สุดของเดือน", f'=_xlfn.MAXIFS({R_AMT},{R_TYPE},"รายจ่าย",{CUR})', BAHT),
       ("หมวดที่ใช้จ่ายมากที่สุด", f'=IFERROR(INDEX({MH["exp"][0]},MATCH(1,{MH["exp"][2]},0)),"-")', None)]
for i, (lab, f_, nf_) in enumerate(ins):
    r = 23 + i
    put(mo, f"B{r}", lab, font(10, False, INK2), b=BORDER, a=al("left", indent=1))
    put(mo, f"C{r}", f_, font(10, True, INK), b=BORDER, a=al("left", indent=1), nf=nf_, merge=f"F{r}")

# KPI cards
sub_cmp = lambda r_: (f'=IF($D${r_}=0,"ไม่มีข้อมูลเดือนก่อน",IF($C${r_}>=$D${r_},"▲ ","▼ ")'
                      f'&TEXT(ABS(($C${r_}-$D${r_})/$D${r_}),"0.0%")&" จากเดือนก่อน")')
card(mo, "B", "C", "💰 รายรับ", "=$C$11", BAHT, sub_cmp(11), None, T["inc"])
card(mo, "D", "F", "💸 รายจ่าย", "=$C$12", BAHT, sub_cmp(12), None, T["exp"])
card(mo, "G", "I", "🏦 เงินออม", "=$C$13", BAHT, sub_cmp(13), None, T["sav"])
card(mo, "J", "L", "💵 คงเหลือ (Cash Flow)", "=$C$14", BAHT,
     '=IF($C$11=0,"","คิดเป็น "&TEXT($C$14/$C$11,"0%")&" ของรายรับ")', None, CASH)
card(mo, "M", "O", "📈 อัตราการออม", "=IF($C$11=0,0,$C$13/$C$11)", "0.0%",
     '="เดือนก่อน "&TEXT(IF($D$11=0,0,$D$13/$D$11),"0.0%")', None, RATE)
card(mo, "P", "Q", "🧾 จำนวนรายการ", f"=COUNTIFS({CUR})", '0" รายการ"', "บันทึกในเดือนนี้", None, CNT)
heights(mo, {5: 24, 6: 38, 7: 22, 8: 14, 9: 26, 10: 24, 15: 8, 16: 24, 21: 8, 22: 22, 26: 10})

# Top 10 + pies
section(mo, "B28", "Q28", "🥧 สัดส่วนตามหมวดหมู่ & Top 10 ของเดือนนี้")
top10(mo, 46, MH, " (เดือนนี้)")
heights(mo, {27: 10, 28: 26, 45: 10, 46: 24, 47: 22})
put(mo, "S39", "ข้อมูลประกอบกราฟ", font(9, True, INK2), merge="AC39")
mon_cd = chart_data_block(mo, 40, MH, [("S", "T"), ("W", "X"), ("AA", "AB")])

MON_W = {**DASH_W}
ch = bar_chart("เดือนนี้ vs เดือนก่อน")
ch.add_data(Reference(mo, min_col=3, max_col=4, min_row=10, max_row=14), titles_from_data=True)
ch.set_categories(Reference(mo, min_col=2, min_row=11, max_row=14))
color_series(ch.series[0], "2A78D6")
color_series(ch.series[1], "B8C4D6")
place(mo, ch, ci("G"), 10, ci("K"), 26, MON_W)

dn = doughnut(mo, "รายรับเดือนนี้ถูกใช้ไปอย่างไร", Reference(mo, min_col=2, min_row=17, max_row=19),
              Reference(mo, min_col=3, min_row=17, max_row=19),
              [T["exp"]["color"], T["sav"]["color"], CASH["color"]])
place(mo, dn, ci("L"), 10, ci("Q"), 26, MON_W)

for key, (a, b) in zip(("inc", "exp", "sav"), [("B", "E"), ("G", "K"), ("M", "Q")]):
    cc, vc, r1, r2 = mon_cd[key]
    pc = pie_labeled(f"สัดส่วน{T[key]['name']}เดือนนี้",
                     Reference(mo, min_col=cc, min_row=r1, max_row=r2),
                     Reference(mo, min_col=vc, min_row=r1, max_row=r2), pie_cols)
    place(mo, pc, ci(a), 29, ci(b), 44, MON_W)
mo.freeze_panes = "A4"
mo.print_area = "A1:Q59"

if PREVIEW:  # หน้าละ 1 แผ่นเวลาแปลงเป็น PDF เพื่อทำภาพตัวอย่าง
    for ws_ in (ws_dash, ws_mon):
        ws_.page_setup.fitToHeight = 1
wb.active = 0
wb.save(OUT)
print("saved", OUT)
